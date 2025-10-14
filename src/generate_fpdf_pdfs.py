#!/usr/bin/env python3
"""
FPDF PDF Generator for Forensic Research

Generates PDFs using FPDF library as a fourth PDF generation source.
This serves as an alternative to LibreOffice/OpenOffice for the forensics research.
"""

import os
from fpdf import FPDF
import docx
import random

def extract_text_from_docx(docx_path):
    """Extract plain text from a Word document."""
    try:
        doc = docx.Document(docx_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text.strip())
        
        return text_content
    except Exception as e:
        print(f"Error reading {docx_path}: {e}")
        return None

def clean_text_for_fpdf(text):
    """Clean text to be compatible with FPDF's latin-1 encoding."""
    # Replace common Unicode characters with ASCII equivalents
    replacements = {
        '\u2013': '-',  # en dash
        '\u2014': '-',  # em dash
        '\u2018': "'",  # left single quote
        '\u2019': "'",  # right single quote
        '\u201c': '"',  # left double quote
        '\u201d': '"',  # right double quote
        '\u2026': '...',  # horizontal ellipsis
        '\u00e9': 'e',  # é
        '\u00e8': 'e',  # è
        '\u00ea': 'e',  # ê
        '\u00eb': 'e',  # ë
        '\u00e0': 'a',  # à
        '\u00e2': 'a',  # â
        '\u00e4': 'a',  # ä
        '\u00f6': 'o',  # ö
        '\u00fc': 'u',  # ü
        '\u00df': 'ss', # ß
        '\u010c': 'C',  # Č
        '\u0259': 'e',  # ə
        '\u20ae': 'R',  # ₮ (tugrik symbol)
    }
    
    cleaned = text
    for unicode_char, ascii_char in replacements.items():
        cleaned = cleaned.replace(unicode_char, ascii_char)
    
    # Remove any remaining non-latin-1 characters
    cleaned = ''.join(c for c in cleaned if ord(c) < 256)
    
    return cleaned

class PDF(FPDF):
    def header(self):
        # Add a simple header
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, 'FPDF Generated Document', 0, 1, 'C')
        self.ln(5)

    def footer(self):
        # Add a simple footer
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

def create_pdf_with_fpdf(text_content, output_path, title):
    """
    Generate PDF using FPDF from text content.
    
    Args:
        text_content (list): List of paragraphs
        output_path (str): Output PDF file path
        title (str): Document title
    """
    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=10)
    
    # Add title
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, title.replace('_', ' '), ln=True, align='C')
    pdf.ln(10)
    
    # Reset font for body
    pdf.set_font("Arial", size=10)
    
    # Add content
    for paragraph_text in text_content:
        if paragraph_text.strip():
            # Clean up text for FPDF
            clean_text = clean_text_for_fpdf(paragraph_text)
            clean_text = clean_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            
            # Split long paragraphs
            words = clean_text.split()
            line = ""
            for word in words:
                if pdf.get_string_width(line + word) < 180:  # Approximate line width
                    line += word + " "
                else:
                    pdf.multi_cell(0, 5, line)
                    line = word + " "
            if line:
                pdf.multi_cell(0, 5, line)
            pdf.ln(3)
    
    # Save PDF
    pdf.output(output_path)

def generate_fpdf_pdfs(docx_dir='wikipedia_docs_expanded', output_dir='fpdf_pdfs', max_files=None):
    """
    Generate FPDF PDFs from existing Word documents.
    
    Args:
        docx_dir (str): Directory containing .docx files
        output_dir (str): Directory to save FPDF-generated PDFs
        max_files (int): Limit number of files to process
    """
    # Create output directory
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # Get list of Word documents
    docx_files = [f for f in os.listdir(docx_dir) if f.endswith('.docx')]
    
    if max_files:
        # Randomly sample files for variety
        docx_files = random.sample(docx_files, min(max_files, len(docx_files)))
    
    print(f"Generating FPDF PDFs for {len(docx_files)} documents...")
    print(f"Input: {docx_dir}")
    print(f"Output: {output_dir}")
    print("=" * 60)
    
    successful = 0
    failed = 0
    
    for i, docx_file in enumerate(docx_files, 1):
        try:
            # Extract text from Word document
            docx_path = os.path.join(docx_dir, docx_file)
            text_content = extract_text_from_docx(docx_path)
            
            if text_content is None:
                print(f"[{i:3d}/{len(docx_files)}] FAILED to read {docx_file}")
                failed += 1
                continue
            
            # Generate PDF
            pdf_filename = docx_file.replace('.docx', '.pdf')
            pdf_path = os.path.join(output_dir, pdf_filename)
            
            title = docx_file.replace('.docx', '').replace('_', ' ')
            create_pdf_with_fpdf(text_content, pdf_path, title)
            
            print(f"[{i:3d}/{len(docx_files)}] {docx_file} -> {pdf_filename}")
            successful += 1
            
        except Exception as e:
            print(f"[{i:3d}/{len(docx_files)}] ERROR processing {docx_file}: {e}")
            failed += 1
    
    print(f"\nFPDF PDF generation complete!")
    print(f"Successful: {successful}")
    print(f"Failed: {failed}")
    
    return successful, failed

def main():
    """Main execution function."""
    print("FPDF PDF Generator for Forensic Research")
    print("=" * 50)
    
    # Check if FPDF is installed
    try:
        from fpdf import FPDF
        print("FPDF available")
    except ImportError:
        print("Error: FPDF not installed. Run: pip install fpdf")
        return
    
    # Check if python-docx is installed
    try:
        import docx
        print("python-docx available")
    except ImportError:
        print("Error: python-docx not installed. Run: pip install python-docx")
        return
    
    print()
    
    # Generate FPDF PDFs (limit to 50 for demonstration)
    successful, failed = generate_fpdf_pdfs(max_files=50)
    
    if successful > 0:
        print("\nNext step: Convert FPDF PDFs to binary images")
        print("Run: python pdf_to_binary_image.py --fpdf-pdfs")

if __name__ == "__main__":
    main()